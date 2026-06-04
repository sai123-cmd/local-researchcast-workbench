from __future__ import annotations

import csv
from pathlib import Path


TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".json", ".csv", ".tsv"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | {".pdf", ".docx", ".xlsx", ".xlsm"}


def supported_document(path: Path) -> bool:
    return path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES


def extract_text(path: Path, max_chars: int = 80_000) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown", ".json"}:
        return path.read_text(encoding="utf-8", errors="ignore")[:max_chars]
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        return _read_csv(path, delimiter, max_chars)
    if suffix == ".docx":
        return _read_docx(path, max_chars)
    if suffix == ".pdf":
        return _read_pdf(path, max_chars)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(path, max_chars)
    return ""


def _read_csv(path: Path, delimiter: str, max_chars: int) -> str:
    lines: list[str] = []
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        for row in reader:
            lines.append(" | ".join(row))
            if sum(len(x) for x in lines) > max_chars:
                break
    return "\n".join(lines)[:max_chars]


def _read_docx(path: Path, max_chars: int) -> str:
    from docx import Document

    doc = Document(str(path))
    chunks: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            chunks.append(text)
        if sum(len(x) for x in chunks) > max_chars:
            break
    return "\n".join(chunks)[:max_chars]


def _read_pdf(path: Path, max_chars: int) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    chunks: list[str] = []
    for page in reader.pages[:30]:
        chunks.append(page.extract_text() or "")
        if sum(len(x) for x in chunks) > max_chars:
            break
    return "\n".join(chunks)[:max_chars]


def _read_xlsx(path: Path, max_chars: int) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True, read_only=True)
    chunks: list[str] = []
    for sheet in wb.worksheets:
        chunks.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(max_row=400, values_only=True):
            values = [str(value) for value in row if value is not None]
            if values:
                chunks.append(" | ".join(values))
            if sum(len(x) for x in chunks) > max_chars:
                return "\n".join(chunks)[:max_chars]
    return "\n".join(chunks)[:max_chars]

